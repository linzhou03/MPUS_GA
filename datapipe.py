"""Data pipeline for SEED-IV/SEED-VII -> SEED-V experiments.

Each movie trial is represented as one graph. The 62 EEG channels are graph
nodes and each node contains 90 temporal positions x 5 DE bands = 450 input
features. The spatial graph itself is learned dynamically by the model, so the
PyG ``edge_index`` stored here only contains self-loops as a batching scaffold.

Normalization is performed independently for every channel-time-band feature.
The strict paper mode applies source-domain statistics to both domains; the
robust mode uses unlabeled per-domain statistics to handle incompatible scales
in the two official pre-extracted DE feature packages.
"""

from __future__ import annotations

import os
import pickle
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Tuple

import numpy as np
import torch
from openpyxl import load_workbook
from scipy.io import loadmat
from torch.utils.data import Dataset
from torch_geometric.data import Data


DEFAULT_DATA_ROOT = Path(
    os.environ.get("CAGA_SGA_DATA_ROOT", "/dataset/gzw/seed_series")
)
NUM_CHANNELS = 62
NUM_BANDS = 5
TIME_STEPS = 90
INPUT_FEATURE_DIM = TIME_STEPS * NUM_BANDS
NUM_CLASSES = 3
TARGET_SUBJECTS = 16
SOURCE_DATASETS = ("seed-vii", "seed-iv")
SOURCE_NAMES = {"seed-vii": "SEED-VII", "seed-iv": "SEED-IV"}
SOURCE_SUBJECT_COUNTS = {"seed-vii": 20, "seed-iv": 15}
SOURCE_TIME_STEPS = {"seed-vii": 90, "seed-iv": 64}

CLASS_NAMES = ("positive", "neutral", "negative")
EMOTION_TO_CLASS = {
    "happy": 0,
    "surprise": 0,
    "neutral": 1,
    "sad": 2,
    "fear": 2,
    "disgust": 2,
    "anger": 2,
}
SEED_V_LABEL_TO_CLASS = {
    0: 2,  # disgust
    1: 2,  # fear
    2: 2,  # sad
    3: 1,  # neutral
    4: 0,  # happy
}
SEED_IV_LABEL_TO_CLASS = {
    0: 1,  # neutral
    1: 2,  # sad
    2: 2,  # fear
    3: 0,  # happy
}
SEED_IV_SESSION_LABELS = (
    (1, 2, 3, 0, 2, 0, 0, 1, 0, 1, 2, 1, 1, 1, 2, 3, 2, 2, 3, 3, 0, 3, 0, 3),
    (2, 1, 3, 0, 0, 2, 0, 2, 3, 3, 2, 3, 2, 0, 1, 1, 2, 1, 0, 3, 0, 1, 3, 1),
    (1, 2, 2, 1, 3, 3, 3, 1, 1, 2, 1, 0, 2, 3, 3, 0, 2, 3, 0, 0, 2, 0, 1, 0),
)


@dataclass
class _SourceCache:
    root: Path
    source_dataset: str
    time_steps: int
    features: np.ndarray
    labels: np.ndarray
    subject_ids: np.ndarray
    trial_ids: np.ndarray
    mean: np.ndarray
    std: np.ndarray
    dataset: "EEGGraphDataset"


_SOURCE_CACHES: Dict[Tuple[Path, str], _SourceCache] = {}


class EEGGraphDataset(Dataset):
    """In-memory EEG trial dataset that emits PyTorch Geometric graphs."""

    def __init__(
        self,
        features: np.ndarray,
        labels: np.ndarray,
        subject_ids: np.ndarray,
        trial_ids: np.ndarray,
    ) -> None:
        if features.ndim != 3 or features.shape[1:] != (
            NUM_CHANNELS,
            INPUT_FEATURE_DIM,
        ):
            raise ValueError(
                "Expected features shaped "
                f"[N, {NUM_CHANNELS}, {INPUT_FEATURE_DIM}], got {features.shape}"
            )
        sample_count = features.shape[0]
        for name, values in {
            "labels": labels,
            "subject_ids": subject_ids,
            "trial_ids": trial_ids,
        }.items():
            if len(values) != sample_count:
                raise ValueError(
                    f"{name} has {len(values)} entries for {sample_count} samples"
                )

        self.features = torch.from_numpy(
            np.ascontiguousarray(features, dtype=np.float32)
        )
        self.labels = torch.from_numpy(np.asarray(labels, dtype=np.int64))
        self.subject_ids = torch.from_numpy(np.asarray(subject_ids, dtype=np.int64))
        self.trial_ids = torch.from_numpy(np.asarray(trial_ids, dtype=np.int64))
        nodes = torch.arange(NUM_CHANNELS, dtype=torch.long)
        self.edge_index = torch.stack((nodes, nodes), dim=0)

    def __len__(self) -> int:
        return self.features.shape[0]

    def __getitem__(self, index: int) -> Data:
        return Data(
            x=self.features[index],
            edge_index=self.edge_index,
            y=self.labels[index].view(1),
            subject_id=self.subject_ids[index].view(1),
            trial_id=self.trial_ids[index].view(1),
        )


def _numeric_stem(path: Path) -> int:
    return int(path.stem.split("_")[0])


def _pad_trial(
    channel_time_band: np.ndarray,
    max_kept_steps: int = TIME_STEPS,
) -> np.ndarray:
    """Convert a [62, T, 5] DE trial into a [62, 450] node matrix."""

    trial = np.asarray(channel_time_band, dtype=np.float32)
    if trial.ndim != 3 or trial.shape[0] != NUM_CHANNELS or trial.shape[2] != NUM_BANDS:
        raise ValueError(
            f"Expected a [{NUM_CHANNELS}, T, {NUM_BANDS}] trial, got {trial.shape}"
        )

    aligned = np.zeros((NUM_CHANNELS, TIME_STEPS, NUM_BANDS), dtype=np.float32)
    if not 1 <= max_kept_steps <= TIME_STEPS:
        raise ValueError(f"max_kept_steps must be within 1..{TIME_STEPS}")
    kept_steps = min(trial.shape[1], max_kept_steps)
    aligned[:, :kept_steps, :] = trial[:, :kept_steps, :]
    return aligned.reshape(NUM_CHANNELS, INPUT_FEATURE_DIM)


def _valid_time_mask(features: np.ndarray) -> np.ndarray:
    """Return [N, 1, T, 1] masks for real DE windows rather than padding."""

    values = features.reshape(-1, NUM_CHANNELS, TIME_STEPS, NUM_BANDS)
    valid = np.any(values != 0, axis=(1, 3))
    return valid[:, None, :, None]


def _fit_valid_feature_stats(features: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
    """Fit each channel-time-band coordinate without including zero padding."""

    values = features.reshape(-1, NUM_CHANNELS, TIME_STEPS, NUM_BANDS).astype(
        np.float64, copy=False
    )
    mask = _valid_time_mask(features)
    counts = mask.sum(axis=0, dtype=np.float64)
    safe_counts = np.maximum(counts, 1.0)
    mean = (values * mask).sum(axis=0) / safe_counts
    variance = (((values - mean[None, ...]) ** 2) * mask).sum(axis=0) / safe_counts
    std = np.sqrt(variance)
    empty = counts == 0
    mean = np.where(empty, 0.0, mean)
    std = np.where(empty, 1.0, np.maximum(std, 1e-5))
    return (
        mean.reshape(NUM_CHANNELS, INPUT_FEATURE_DIM).astype(np.float32),
        std.reshape(NUM_CHANNELS, INPUT_FEATURE_DIM).astype(np.float32),
    )


def _normalize_valid_features(
    features: np.ndarray,
    mean: np.ndarray,
    std: np.ndarray,
) -> np.ndarray:
    """Normalize real DE windows while preserving padding as exact zeros."""

    normalized = (
        (features - mean[None, :, :]) / std[None, :, :]
    ).astype(np.float32, copy=False)
    normalized_view = normalized.reshape(
        -1, NUM_CHANNELS, TIME_STEPS, NUM_BANDS
    )
    normalized_view *= _valid_time_mask(features)
    return normalized


def _seed_vii_labels(root: Path) -> np.ndarray:
    label_file = root / "labels/SEED_VII/emotion_label_and_stimuli_order.xlsx"
    if not label_file.is_file():
        raise FileNotFoundError(f"SEED-VII label file not found: {label_file}")

    workbook = load_workbook(label_file, read_only=True, data_only=True)
    sheet = workbook.active
    labels = []
    for row_number in (2, 4, 6, 8):
        row = list(
            sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True)
        )[0]
        emotions = [value for value in row[1:] if value is not None]
        if len(emotions) != 20:
            raise ValueError(
                f"Expected 20 SEED-VII labels in row {row_number}, got {len(emotions)}"
            )
        for emotion in emotions:
            key = str(emotion).strip().lower()
            if key not in EMOTION_TO_CLASS:
                raise ValueError(f"Unknown SEED-VII emotion label: {emotion!r}")
            labels.append(EMOTION_TO_CLASS[key])
    workbook.close()

    if len(labels) != 80:
        raise ValueError(f"Expected 80 SEED-VII trial labels, got {len(labels)}")
    return np.asarray(labels, dtype=np.int64)


def _load_seed_vii(root: Path) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    feature_dir = root / "feature/seed_vii/EEG_features"
    files = sorted(feature_dir.glob("*.mat"), key=_numeric_stem)
    expected_subjects = SOURCE_SUBJECT_COUNTS["seed-vii"]
    if len(files) != expected_subjects:
        raise FileNotFoundError(
            f"Expected {expected_subjects} SEED-VII subject files in {feature_dir}, "
            f"found {len(files)}"
        )

    labels_per_subject = _seed_vii_labels(root)
    features = []
    labels = []
    subject_ids = []
    trial_ids = []
    variable_names = [f"de_LDS_{trial_id}" for trial_id in range(1, 81)]

    for expected_subject, path in enumerate(files, start=1):
        subject_id = _numeric_stem(path)
        if subject_id != expected_subject:
            raise ValueError(
                f"Unexpected SEED-VII subject ordering: expected {expected_subject}, "
                f"found {subject_id} in {path.name}"
            )
        mat = loadmat(path, variable_names=variable_names)
        for trial_id, label in enumerate(labels_per_subject, start=1):
            key = f"de_LDS_{trial_id}"
            if key not in mat:
                raise KeyError(f"Missing {key} in {path}")
            raw = np.asarray(mat[key])
            if raw.ndim != 3 or raw.shape[1:] != (NUM_BANDS, NUM_CHANNELS):
                raise ValueError(f"Unexpected {key} shape in {path}: {raw.shape}")
            features.append(_pad_trial(raw.transpose(2, 0, 1)))
            labels.append(label)
            subject_ids.append(subject_id)
            trial_ids.append(trial_id)

    return (
        np.stack(features),
        np.asarray(labels, dtype=np.int64),
        np.asarray(subject_ids, dtype=np.int64),
        np.asarray(trial_ids, dtype=np.int64),
    )


def _load_seed_iv(root: Path) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """Load the official smoothed DE features from all SEED-IV sessions."""

    feature_root = root / "feature/seed_iv/eeg_feature_smooth"
    expected_subjects = SOURCE_SUBJECT_COUNTS["seed-iv"]
    variable_names = [f"de_LDS{trial_id}" for trial_id in range(1, 25)]
    features = []
    labels = []
    subject_ids = []
    trial_ids = []

    for session_id, session_labels in enumerate(SEED_IV_SESSION_LABELS, start=1):
        session_dir = feature_root / str(session_id)
        files = sorted(session_dir.glob("*.mat"), key=_numeric_stem)
        if len(files) != expected_subjects:
            raise FileNotFoundError(
                f"Expected {expected_subjects} SEED-IV files in {session_dir}, "
                f"found {len(files)}"
            )

        for expected_subject, path in enumerate(files, start=1):
            subject_id = _numeric_stem(path)
            if subject_id != expected_subject:
                raise ValueError(
                    "Unexpected SEED-IV subject ordering: "
                    f"expected {expected_subject}, found {subject_id} in {path.name}"
                )
            mat = loadmat(path, variable_names=variable_names)
            for session_trial, original_label in enumerate(session_labels, start=1):
                key = f"de_LDS{session_trial}"
                if key not in mat:
                    raise KeyError(f"Missing {key} in {path}")
                raw = np.asarray(mat[key])
                if raw.ndim != 3 or raw.shape[0] != NUM_CHANNELS or raw.shape[2] != NUM_BANDS:
                    raise ValueError(f"Unexpected {key} shape in {path}: {raw.shape}")
                features.append(
                    _pad_trial(
                        raw,
                        max_kept_steps=SOURCE_TIME_STEPS["seed-iv"],
                    )
                )
                labels.append(SEED_IV_LABEL_TO_CLASS[int(original_label)])
                subject_ids.append(subject_id)
                trial_ids.append((session_id - 1) * 24 + session_trial)

    return (
        np.stack(features),
        np.asarray(labels, dtype=np.int64),
        np.asarray(subject_ids, dtype=np.int64),
        np.asarray(trial_ids, dtype=np.int64),
    )


def _load_seed_v_subject(
    root: Path,
    subject_id: int,
    max_kept_steps: int = TIME_STEPS,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    feature_file = root / f"feature/seed_v/EEG_DE_features/{subject_id}_123.npz"
    if not feature_file.is_file():
        raise FileNotFoundError(f"SEED-V subject file not found: {feature_file}")

    with np.load(feature_file, allow_pickle=True) as archive:
        if "data" not in archive or "label" not in archive:
            raise KeyError(f"Expected data and label entries in {feature_file}")
        trial_data: Dict[int, np.ndarray] = pickle.loads(archive["data"].item())
        trial_labels: Dict[int, np.ndarray] = pickle.loads(archive["label"].item())

    if sorted(trial_data) != list(range(45)) or sorted(trial_labels) != list(range(45)):
        raise ValueError(f"Expected SEED-V trial keys 0..44 in {feature_file}")

    features = []
    labels = []
    trial_ids = []
    for trial_key in range(45):
        raw = np.asarray(trial_data[trial_key], dtype=np.float32)
        if raw.ndim != 2 or raw.shape[1] != NUM_CHANNELS * NUM_BANDS:
            raise ValueError(
                f"Unexpected SEED-V trial {trial_key} shape in {feature_file}: {raw.shape}"
            )

        # SEED-V stores each time step as 62 consecutive channel blocks, with
        # five band values per channel: [T, 62 * 5].
        channel_time_band = raw.reshape(
            raw.shape[0], NUM_CHANNELS, NUM_BANDS
        ).transpose(1, 0, 2)
        features.append(
            _pad_trial(channel_time_band, max_kept_steps=max_kept_steps)
        )

        original_labels = np.asarray(trial_labels[trial_key]).reshape(-1)
        if original_labels.size == 0 or not np.all(
            original_labels == original_labels[0]
        ):
            raise ValueError(
                f"Inconsistent labels for SEED-V trial {trial_key} in {feature_file}"
            )
        original_label = int(original_labels[0])
        if original_label not in SEED_V_LABEL_TO_CLASS:
            raise ValueError(f"Unknown SEED-V label {original_label} in {feature_file}")
        labels.append(SEED_V_LABEL_TO_CLASS[original_label])
        trial_ids.append(trial_key + 1)

    sample_count = len(features)
    return (
        np.stack(features),
        np.asarray(labels, dtype=np.int64),
        np.full(sample_count, subject_id, dtype=np.int64),
        np.asarray(trial_ids, dtype=np.int64),
    )


def _ensure_source_loaded(root: Path, source_dataset: str) -> _SourceCache:
    root = root.expanduser().resolve()
    if source_dataset not in SOURCE_DATASETS:
        raise ValueError(
            f"source_dataset must be one of {SOURCE_DATASETS}, got {source_dataset!r}"
        )
    cache_key = (root, source_dataset)
    if cache_key in _SOURCE_CACHES:
        return _SOURCE_CACHES[cache_key]

    loader = _load_seed_vii if source_dataset == "seed-vii" else _load_seed_iv
    features, labels, subject_ids, trial_ids = loader(root)
    # A channel-time-band coordinate is one feature dimension in Eq. (2).
    # Pooling channels here would incorrectly force different electrodes to
    # share normalization statistics and erase electrode-specific structure.
    mean, std = _fit_valid_feature_stats(features)
    normalized = _normalize_valid_features(features, mean, std)
    dataset = EEGGraphDataset(normalized, labels, subject_ids, trial_ids)
    cache = _SourceCache(
        root=root,
        source_dataset=source_dataset,
        time_steps=SOURCE_TIME_STEPS[source_dataset],
        features=normalized,
        labels=labels,
        subject_ids=subject_ids,
        trial_ids=trial_ids,
        mean=mean,
        std=std,
        dataset=dataset,
    )
    _SOURCE_CACHES[cache_key] = cache
    return cache


def _class_counts(labels: np.ndarray) -> Dict[str, int]:
    counts = Counter(int(label) for label in labels)
    return {CLASS_NAMES[index]: counts.get(index, 0) for index in range(NUM_CLASSES)}


def create_cross_dataset_setup(
    data_root: str | os.PathLike[str] = DEFAULT_DATA_ROOT,
    source_dataset: str = "seed-vii",
) -> dict:
    """Validate and load the source domain, returning a concise data summary."""

    root = Path(data_root)
    source = _ensure_source_loaded(root, source_dataset)
    target_dir = source.root / "feature/seed_v/EEG_DE_features"
    target_files = sorted(target_dir.glob("*_123.npz"), key=_numeric_stem)
    if len(target_files) != TARGET_SUBJECTS:
        raise FileNotFoundError(
            f"Expected {TARGET_SUBJECTS} SEED-V subject files in {target_dir}, "
            f"found {len(target_files)}"
        )

    summary = {
        "data_root": str(source.root),
        "source": SOURCE_NAMES[source_dataset],
        "source_subjects": SOURCE_SUBJECT_COUNTS[source_dataset],
        "source_trials": len(source.dataset),
        "source_class_counts": _class_counts(source.labels),
        "target": "SEED-V",
        "target_subjects": TARGET_SUBJECTS,
        "trials_per_target_subject": 45,
        "feature_shape": (NUM_CHANNELS, INPUT_FEATURE_DIM),
        "effective_time_steps": source.time_steps,
        "normalization": "source-only per-channel-feature z-score",
    }
    print("Dataset setup:", summary)
    return summary


def load_cross_dataset_fold(
    target_subject_index: int,
    data_root: str | os.PathLike[str] = DEFAULT_DATA_ROOT,
    target_normalization: str = "source",
    source_dataset: str = "seed-vii",
) -> Tuple[EEGGraphDataset, EEGGraphDataset]:
    """Return all labeled source trials and one unlabeled target-subject fold.

    ``target_subject_index`` is zero-based to remain compatible with the
    original training script. Labels remain attached to target graphs solely
    for final evaluation; the training loop must not read them.
    """

    if not 0 <= target_subject_index < TARGET_SUBJECTS:
        raise ValueError(
            f"target_subject_index must be in [0, {TARGET_SUBJECTS - 1}], "
            f"got {target_subject_index}"
        )
    if target_normalization not in {"source", "domain"}:
        raise ValueError(
            "target_normalization must be either 'source' or 'domain', "
            f"got {target_normalization!r}"
        )

    source = _ensure_source_loaded(Path(data_root), source_dataset)
    target_subject_id = target_subject_index + 1
    features, labels, subject_ids, trial_ids = _load_seed_v_subject(
        source.root,
        target_subject_id,
        max_kept_steps=source.time_steps,
    )
    if target_normalization == "domain":
        # This diagnostic mode uses only unlabeled target features. It is
        # useful when the official pre-extracted DE files have incompatible
        # numerical scales, while ``source`` retains the paper's strict Eq. (2)
        # protocol.
        target_mean, target_std = _fit_valid_feature_stats(features)
        normalized = _normalize_valid_features(features, target_mean, target_std)
    else:
        normalized = _normalize_valid_features(features, source.mean, source.std)
    target_dataset = EEGGraphDataset(normalized, labels, subject_ids, trial_ids)
    print(
        f"Target subject {target_subject_id:02d}: {len(target_dataset)} trials, "
        f"class counts {_class_counts(labels)}"
    )
    return source.dataset, target_dataset
