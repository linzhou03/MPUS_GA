"""Dataset-specific trial boundaries and emotion mappings."""

from __future__ import annotations

import csv
import datetime as dt
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


EMOTION_TO_ORIGINAL_SEED_VII = {
    "disgust": 0,
    "fear": 1,
    "sad": 2,
    "neutral": 3,
    "happy": 4,
    "anger": 5,
    "surprise": 6,
}

EMOTION_TO_ORIGINAL_SEED_V = {
    "disgust": 0,
    "fear": 1,
    "sad": 2,
    "neutral": 3,
    "happy": 4,
}

EMOTION_TO_ORIGINAL_SEED_IV = {
    "neutral": 0,
    "sad": 1,
    "fear": 2,
    "happy": 3,
}

EMOTION_TO_THREE_CLASS = {
    "happy": 0,
    "surprise": 0,
    "neutral": 1,
    "disgust": 2,
    "fear": 2,
    "sad": 2,
    "anger": 2,
}

SEED_V_START_SECONDS = {
    1: [30, 132, 287, 555, 773, 982, 1271, 1628, 1730, 2025, 2227, 2435, 2667, 2932, 3204],
    2: [30, 299, 548, 646, 836, 1000, 1091, 1392, 1657, 1809, 1966, 2186, 2333, 2490, 2741],
    3: [30, 353, 478, 674, 825, 908, 1200, 1346, 1451, 1711, 2055, 2307, 2457, 2726, 2888],
}

SEED_V_END_SECONDS = {
    1: [102, 228, 524, 742, 920, 1240, 1568, 1697, 1994, 2166, 2401, 2607, 2901, 3172, 3359],
    2: [267, 488, 614, 773, 967, 1059, 1331, 1622, 1777, 1908, 2153, 2302, 2428, 2709, 2817],
    3: [321, 418, 643, 764, 877, 1147, 1284, 1418, 1679, 1996, 2275, 2425, 2664, 2857, 3066],
}

SEED_V_EMOTIONS = {
    1: [
        "happy", "fear", "neutral", "sad", "disgust",
        "happy", "fear", "neutral", "sad", "disgust",
        "happy", "fear", "neutral", "sad", "disgust",
    ],
    2: [
        "sad", "fear", "neutral", "disgust", "happy",
        "happy", "disgust", "neutral", "sad", "fear",
        "neutral", "happy", "fear", "sad", "disgust",
    ],
    3: [
        "sad", "fear", "neutral", "disgust", "happy",
        "happy", "disgust", "neutral", "sad", "fear",
        "neutral", "happy", "fear", "sad", "disgust",
    ],
}

SEED_IV_SESSION_LABELS = {
    1: (1, 2, 3, 0, 2, 0, 0, 1, 0, 1, 2, 1, 1, 1, 2, 3, 2, 2, 3, 3, 0, 3, 0, 3),
    2: (2, 1, 3, 0, 0, 2, 0, 2, 3, 3, 2, 3, 2, 0, 1, 1, 2, 1, 0, 3, 0, 1, 3, 1),
    3: (1, 2, 2, 1, 3, 3, 3, 1, 1, 2, 1, 0, 2, 3, 3, 0, 2, 3, 0, 0, 2, 0, 1, 0),
}

SEED_IV_LABEL_TO_EMOTION = {
    0: "neutral",
    1: "sad",
    2: "fear",
    3: "happy",
}

SEED_IV_EMOTIONS = {
    session: tuple(SEED_IV_LABEL_TO_EMOTION[label] for label in labels)
    for session, labels in SEED_IV_SESSION_LABELS.items()
}

SEED_VII_SPECIAL_TRIGGER_START = {
    "14_20221015_1.cnt": "14:25:34",
    "9_20221111_3.cnt": "14:01:27",
}


def parse_seed_v_filename(path: Path) -> tuple[int, int]:
    match = re.fullmatch(r"(\d+)_(\d)_(\d{8})\.cnt", path.name)
    if not match:
        raise ValueError(f"Unexpected SEED-V filename: {path.name}")
    return int(match.group(1)), int(match.group(2))


def parse_seed_vii_filename(path: Path) -> tuple[int, int]:
    match = re.fullmatch(r"(\d+)_(\d{8})_(\d)\.cnt", path.name)
    if not match:
        raise ValueError(f"Unexpected SEED-VII filename: {path.name}")
    return int(match.group(1)), int(match.group(3))


def parse_seed_iv_filename(path: Path) -> tuple[int, int]:
    match = re.fullmatch(r"(\d+)_(\d{8})\.mat", path.name)
    if not match:
        raise ValueError(f"Unexpected SEED-IV filename: {path.name}")
    try:
        session = int(path.parent.name)
    except ValueError as exc:
        raise ValueError(
            f"SEED-IV file must be inside session directory 1, 2, or 3: {path}"
        ) from exc
    if session not in SEED_IV_SESSION_LABELS:
        raise ValueError(
            f"SEED-IV file must be inside session directory 1, 2, or 3: {path}"
        )
    return int(match.group(1)), session


def load_channel_names(channel_file: Path) -> tuple[str, ...]:
    if not channel_file.is_file():
        raise FileNotFoundError(f"Channel order file not found: {channel_file}")
    workbook = load_workbook(channel_file, read_only=True, data_only=True)
    try:
        channels = tuple(
            str(row[0]).strip()
            for row in workbook.active.iter_rows(values_only=True)
            if row[0] is not None and str(row[0]).strip()
        )
    finally:
        workbook.close()
    if len(channels) != 62 or len(set(channels)) != 62:
        raise ValueError(
            f"Expected 62 unique channel names in {channel_file}, got {len(channels)}"
        )
    return channels


def load_seed_vii_emotions(label_file: Path) -> dict[int, str]:
    if not label_file.is_file():
        raise FileNotFoundError(f"SEED-VII label file not found: {label_file}")
    workbook = load_workbook(label_file, read_only=True, data_only=True)
    sheet = workbook.active
    emotions: list[str] = []
    for row_number in (2, 4, 6, 8):
        row = next(
            sheet.iter_rows(
                min_row=row_number,
                max_row=row_number,
                values_only=True,
            )
        )
        session_emotions = [
            str(value).strip().lower() for value in row[1:] if value is not None
        ]
        if len(session_emotions) != 20:
            raise ValueError(
                f"Expected 20 SEED-VII labels in Excel row {row_number}, "
                f"got {len(session_emotions)}"
            )
        unknown = sorted(set(session_emotions) - set(EMOTION_TO_ORIGINAL_SEED_VII))
        if unknown:
            raise ValueError(f"Unknown SEED-VII emotions: {unknown}")
        emotions.extend(session_emotions)
    workbook.close()
    if len(emotions) != 80:
        raise ValueError(f"Expected 80 SEED-VII labels, got {len(emotions)}")
    return {index + 1: emotion for index, emotion in enumerate(emotions)}


def read_special_trigger_samples(
    save_info: Path,
    raw_name: str,
    sfreq: float,
) -> list[int]:
    if raw_name not in SEED_VII_SPECIAL_TRIGGER_START:
        raise ValueError(f"No special trigger configuration for {raw_name}")
    stem = Path(raw_name).stem
    csv_name = f"{stem}_trigger_info.csv"
    rows: list[list[str]]
    if save_info.is_dir():
        csv_path = save_info / csv_name
        if not csv_path.is_file():
            csv_path = save_info / "save_info" / csv_name
        with csv_path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
    elif save_info.suffix.lower() == ".zip":
        with zipfile.ZipFile(save_info) as archive:
            candidates = (csv_name, f"save_info/{csv_name}")
            member = next((name for name in candidates if name in archive.namelist()), None)
            if member is None:
                raise FileNotFoundError(
                    f"Could not find {csv_name} in {save_info}"
                )
            with archive.open(member) as handle:
                rows = list(
                    csv.reader(
                        line.decode("utf-8-sig").strip() for line in handle
                    )
                )
    else:
        raise FileNotFoundError(
            f"SEED-VII save_info must be a directory or ZIP: {save_info}"
        )

    start = dt.datetime.strptime(
        SEED_VII_SPECIAL_TRIGGER_START[raw_name], "%H:%M:%S"
    )
    samples = []
    for row in rows:
        event_time = dt.datetime.strptime(row[1].split(" ")[-1], "%H:%M:%S.%f")
        samples.append(int(round((event_time - start).total_seconds() * sfreq)))
    return samples
